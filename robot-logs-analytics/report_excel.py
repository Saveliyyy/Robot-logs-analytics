"""
Excel-отчёт для экспресс-анализа: данные витрины + сводка на живых формулах.

Лист «Сводка» считается формулами СУММЕСЛИМН / СРЗНАЧЕСЛИМН / МАКСЕСЛИМН,
поэтому достаточно поменять даты периода в жёлтых ячейках, и всё пересчитается.

Запуск:  python report_excel.py
Результат: output/robot_report.xlsx
"""
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = Path(__file__).parent
DB_PATH = BASE / "data" / "warehouse.db"
OUT = BASE / "output" / "robot_report.xlsx"

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="D0D7DE")


def style_header(ws, row: int, ncols: int, start_col: int = 1) -> None:
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_df(ws, df: pd.DataFrame, table_name: str, formats: dict | None = None) -> None:
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    style_header(ws, 1, len(df.columns))
    ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)
    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = max(12, len(str(col)) + 4)
        fmt = (formats or {}).get(col)
        for r in range(2, len(df) + 2):
            cell = ws[f"{letter}{r}"]
            cell.font = Font(name=FONT)
            if fmt:
                cell.number_format = fmt
    ws.freeze_panes = "A2"


def main():
    with sqlite3.connect(DB_PATH) as conn:
        daily = pd.read_sql("SELECT event_date, robot_id, events, errors, warnings, "
                            "error_rate_pct, avg_temp, max_temp, min_battery, p95_latency_ms "
                            "FROM daily_robot_metrics ORDER BY event_date, robot_id", conn)
        etl = pd.read_sql("SELECT event_date, raw_lines, duplicates, malformed, bad_temp, "
                          "missing_battery, clean_rows FROM etl_runs ORDER BY event_date", conn)
    daily["event_date"] = pd.to_datetime(daily["event_date"])
    daily["error_rate_pct"] = daily["error_rate_pct"] / 100     # доля, формат 0.00%
    daily = daily.rename(columns={
        "event_date": "Дата", "robot_id": "Робот", "events": "События", "errors": "Ошибки",
        "warnings": "Предупреждения", "error_rate_pct": "Доля ошибок", "avg_temp": "Ср. темп., °C",
        "max_temp": "Макс. темп., °C", "min_battery": "Мин. заряд, %",
        "p95_latency_ms": "p95 задержки, мс"})
    etl["event_date"] = pd.to_datetime(etl["event_date"])
    etl = etl.rename(columns={
        "event_date": "Дата", "raw_lines": "Строк на входе", "duplicates": "Дубли",
        "malformed": "Битые строки", "bad_temp": "Мусор в temp",
        "missing_battery": "Пропуски заряда", "clean_rows": "Чистых строк"})

    wb = Workbook()

    # ---------- Лист «Данные» ----------
    ws_data = wb.active
    ws_data.title = "Данные"
    write_df(ws_data, daily, "DailyMetrics", {
        "Дата": "dd.mm.yyyy", "События": "#,##0", "Ошибки": "#,##0",
        "Доля ошибок": "0.00%", "Ср. темп., °C": "0.0", "Макс. темп., °C": "0.0",
        "Мин. заряд, %": "0.0", "p95 задержки, мс": "0"})
    n = len(daily) + 1
    rng = {name: f"Данные!${get_column_letter(i)}$2:${get_column_letter(i)}${n}"
           for i, name in enumerate(daily.columns, start=1)}

    # ---------- Лист «Сводка» ----------
    ws = wb.create_sheet("Сводка", 0)
    ws["A1"] = "Сводка по роботам за период"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A3"], ws["A4"] = "Начало периода", "Конец периода"
    ws["B3"], ws["B4"] = daily["Дата"].min().to_pydatetime(), daily["Дата"].max().to_pydatetime()
    for c in ("B3", "B4"):
        ws[c].fill = INPUT_FILL
        ws[c].number_format = "dd.mm.yyyy"
        ws[c].font = Font(name=FONT, color="0000FF")
    ws["C3"] = "← меняйте даты в жёлтых ячейках, сводка пересчитается"
    ws["C3"].font = Font(name=FONT, italic=True, color="666666")
    ws["B3"].comment = Comment("Ввод пользователя. По умолчанию — весь период данных.", "report")
    for c in ("A3", "A4"):
        ws[c].font = Font(name=FONT, bold=True)

    headers = ["Робот", "События", "Ошибки", "Доля ошибок", "Ср. темп., °C",
               "Макс. темп., °C", "Статус"]
    top = 6
    for i, h in enumerate(headers, start=1):
        ws.cell(row=top, column=i, value=h)
    style_header(ws, top, len(headers))

    period = (f'{rng["Дата"]},">="&$B$3,{rng["Дата"]},"<="&$B$4')
    robots = sorted(daily["Робот"].unique())
    first = top + 1
    for k, robot in enumerate(robots):
        r = first + k
        ws[f"A{r}"] = robot
        ws[f"B{r}"] = f'=SUMIFS({rng["События"]},{rng["Робот"]},$A{r},{period})'
        ws[f"C{r}"] = f'=SUMIFS({rng["Ошибки"]},{rng["Робот"]},$A{r},{period})'
        ws[f"D{r}"] = f"=IFERROR(C{r}/B{r},0)"
        ws[f"E{r}"] = f'=IFERROR(AVERAGEIFS({rng["Ср. темп., °C"]},{rng["Робот"]},$A{r},{period}),0)'
        ws[f"F{r}"] = f'=_xlfn.MAXIFS({rng["Макс. темп., °C"]},{rng["Робот"]},$A{r},{period})'
        ws[f"G{r}"] = f'=IF(D{r}>2*$D${first + len(robots)},"Проверить","Норма")'
    total = first + len(robots)
    ws[f"A{total}"] = "Весь парк"
    ws[f"B{total}"] = f"=SUM(B{first}:B{total - 1})"
    ws[f"C{total}"] = f"=SUM(C{first}:C{total - 1})"
    ws[f"D{total}"] = f"=IFERROR(C{total}/B{total},0)"
    ws[f"E{total}"] = f"=AVERAGE(E{first}:E{total - 1})"
    ws[f"F{total}"] = f"=MAX(F{first}:F{total - 1})"
    ws[f"G{total}"] = "Статус «Проверить»: доля ошибок робота > 2× средней по парку"
    ws[f"G{total}"].font = Font(name=FONT, italic=True, color="666666")

    for r in range(first, total + 1):
        for c in range(1, 7):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, bold=(r == total))
            cell.border = Border(bottom=THIN)
        ws[f"B{r}"].number_format = "#,##0"
        ws[f"C{r}"].number_format = "#,##0"
        ws[f"D{r}"].number_format = "0.00%"
        ws[f"E{r}"].number_format = "0.0"
        ws[f"F{r}"].number_format = "0.0"
    for r in range(first, total):
        ws[f"G{r}"].font = Font(name=FONT, bold=True)
    ws.conditional_formatting.add(
        f"D{first}:D{total - 1}",
        ColorScaleRule(start_type="min", start_color="E6F4EA",
                       end_type="max", end_color="F4C7C3"))

    # Быстрый поиск: ошибки робота в конкретный день (ИНДЕКС+ПОИСКПОЗ по двум условиям)
    q = total + 3
    ws[f"A{q}"] = "Быстрый поиск"
    ws[f"A{q}"].font = Font(name=FONT, bold=True, size=12)
    ws[f"A{q + 1}"], ws[f"B{q + 1}"] = "Робот", robots[-1]
    ws[f"A{q + 2}"], ws[f"B{q + 2}"] = "Дата", daily["Дата"].max().to_pydatetime()
    ws[f"B{q + 2}"].number_format = "dd.mm.yyyy"
    for c in (f"B{q + 1}", f"B{q + 2}"):
        ws[c].fill = INPUT_FILL
        ws[c].font = Font(name=FONT, color="0000FF")
    ws[f"A{q + 3}"] = "Ошибок в этот день"
    ws[f"B{q + 3}"] = (f'=IFERROR(INDEX({rng["Ошибки"]},MATCH(1,INDEX(({rng["Робот"]}=B{q + 1})'
                       f'*({rng["Дата"]}=B{q + 2}),0),0)),"нет данных")')
    ws[f"A{q + 4}"] = "Макс. температура"
    ws[f"B{q + 4}"] = (f'=IFERROR(INDEX({rng["Макс. темп., °C"]},MATCH(1,INDEX(({rng["Робот"]}=B{q + 1})'
                       f'*({rng["Дата"]}=B{q + 2}),0),0)),"нет данных")')
    for r in range(q + 1, q + 5):
        ws[f"A{r}"].font = Font(name=FONT, bold=True)
    ws[f"B{q + 3}"].font = Font(name=FONT)
    ws[f"B{q + 4}"].font = Font(name=FONT)
    ws[f"B{q + 4}"].number_format = "0.0"

    for col, w in zip("ABCDEFG", (22, 13, 11, 13, 14, 15, 14)):
        ws.column_dimensions[col].width = w

    chart = BarChart()
    chart.type = "col"
    chart.title = "Доля ошибок по роботам"
    chart.y_axis.numFmt = "0.0%"
    chart.y_axis.majorGridlines = None
    chart.legend = None
    chart.add_data(Reference(ws, min_col=4, min_row=top, max_row=total - 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=total - 1))
    chart.height, chart.width = 7, 13
    ws.add_chart(chart, "I3")

    # ---------- Лист «Прогноз» ----------
    fc_path = BASE / "output" / "forecast_test.csv"
    if fc_path.exists():
        fc = pd.read_csv(fc_path, parse_dates=["event_date"])[
            ["event_date", "robot_id", "errors", "pred_model"]]
        fc.columns = ["Дата", "Робот", "Факт ошибок", "Прогноз модели"]
        ws_fc = wb.create_sheet("Прогноз")
        write_df(ws_fc, fc, "ForecastTest", {"Дата": "dd.mm.yyyy", "Прогноз модели": "0.0"})
        ws_fc["E1"] = "Абс. ошибка"
        style_header(ws_fc, 1, 1, start_col=5)
        for r in range(2, len(fc) + 2):
            ws_fc[f"E{r}"] = f"=ABS(C{r}-D{r})"
            ws_fc[f"E{r}"].number_format = "0.0"
            ws_fc[f"E{r}"].font = Font(name=FONT)
        ws_fc["G1"], ws_fc["H1"] = "MAE модели, ошибок/день", f"=AVERAGE(E2:E{len(fc) + 1})"
        ws_fc["G1"].font = Font(name=FONT, bold=True)
        ws_fc["H1"].number_format = "0.00"
        ws_fc.column_dimensions["E"].width = 13
        ws_fc.column_dimensions["G"].width = 26

    # ---------- Лист «Качество ETL» ----------
    ws_q = wb.create_sheet("Качество ETL")
    write_df(ws_q, etl, "EtlQuality", {"Дата": "dd.mm.yyyy", "Строк на входе": "#,##0",
                                       "Чистых строк": "#,##0"})
    last = len(etl) + 1
    ws_q["I1"], ws_q["J1"] = "Доля отброшенных строк", f"=1-SUM(G2:G{last})/SUM(B2:B{last})"
    ws_q["I1"].font = Font(name=FONT, bold=True)
    ws_q["J1"].number_format = "0.00%"
    ws_q.column_dimensions["I"].width = 24

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"Отчёт сохранён: {OUT}")


if __name__ == "__main__":
    main()
